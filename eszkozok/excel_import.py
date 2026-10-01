#!/usr/bin/env python3
"""A szerkesztett Excel (szerkesztes/kerdesek.xlsx) visszatöltése az adatok/*.json fájlokba.

Futtatás a projekt gyökeréből:
    python3 eszkozok/excel_import.py [útvonal/a/kerdesek.xlsx]

Előbb mindent beolvas és ellenőriz; csak akkor ír, ha nincs hiba. Utána
érdemes lefuttatni az eszkozok/ellenorzes.py-t is."""
import json, pathlib, sys
from openpyxl import load_workbook

GYOKER = pathlib.Path(__file__).resolve().parent.parent
PALYAK = [("politikai", "Politikai"), ("vallalati", "Vállalati"), ("katonai", "Katonai")]
PONTKULCSOK = ["kovetkezetesseg", "atlathatosag", "dontesi_minoseg", "felelossegvallalas", "erintettek_bevonasa"]
PONTFEJ = ["Következetesség", "Átláthatóság", "Döntési minőség", "Felelősségvállalás", "Érintettek bevonása"]
MEZOK = {  # fejléc eleje → belső név
    "Sz.": "id", "Jelenet címe": "cim", "Helyzet": "szituacio_leirasa", "Jel": "jel",
    "Válasz": "opcio_leirasa", "Fogalom": "fogalom", "Kiértékelés": "kiertekeles",
}

hibak = []

def szoveg(v):
    """Cellaérték tisztítva: Excelből jöhet nem törő szóköz, sorvégi szóköz, Windows-sortörés."""
    if v is None: return ""
    return str(v).replace(" ", " ").replace("\r\n", "\n").strip()

def egesz(v, hol):
    if isinstance(v, bool) or v is None or v == "":
        hibak.append(f"{hol}: hiányzó pont"); return 0
    try:
        f = float(str(v).replace("−", "-").replace(",", "."))
    except ValueError:
        hibak.append(f"{hol}: nem szám ({v!r})"); return 0
    if f != int(f) or not -3 <= f <= 3:
        hibak.append(f"{hol}: {v!r} — egész szám kell −3 és +3 között")
    return int(f)

def oszlopok(ws):
    fej = [szoveg(c.value) for c in ws[1]]
    hely = {}
    for i, f in enumerate(fej):
        for eleje, nev in MEZOK.items():
            # a rövid fejléceknél pontos egyezés kell: a „Jel” különben a „Jelenet címe” elejére is illene
            illik = f == eleje if len(eleje) <= 3 else f.startswith(eleje)
            if illik and nev not in hely: hely[nev] = i
        for kulcs, pf in zip(PONTKULCSOK, PONTFEJ):
            if f == pf: hely[kulcs] = i
    hiany = [n for n in list(MEZOK.values()) + PONTKULCSOK if n not in hely]
    if hiany: hibak.append(f"{ws.title}: hiányzó oszlop(ok): {', '.join(hiany)}")
    return hely

def palya(ws, regi):
    h = oszlopok(ws)
    if hibak: return None
    szituaciok, aktualis = [], None
    for r, sor in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if all(v is None or szoveg(v) == "" for v in sor): continue
        hol = f"{ws.title} {r}. sor"
        if sor[h["id"]] not in (None, ""):          # összevont cella: csak a csoport első sorában van érték
            aktualis = {"id": int(sor[h["id"]]), "cim": szoveg(sor[h["cim"]]),
                        "szituacio_leirasa": szoveg(sor[h["szituacio_leirasa"]]), "opciok": []}
            szituaciok.append(aktualis)
        if aktualis is None:
            hibak.append(f"{hol}: válasz helyzet nélkül"); continue
        aktualis["opciok"].append({
            "jel": szoveg(sor[h["jel"]]),
            "opcio_leirasa": szoveg(sor[h["opcio_leirasa"]]),
            "fogalom": szoveg(sor[h["fogalom"]]),
            "kiertekeles": szoveg(sor[h["kiertekeles"]]),
            "pontok": {k: egesz(sor[h[k]], f"{hol}, {pf}") for k, pf in zip(PONTKULCSOK, PONTFEJ)},
        })
    for s in szituaciok:
        hol = f"{ws.title} {s['id']}. helyzet"
        if not s["cim"] or not s["szituacio_leirasa"]: hibak.append(f"{hol}: üres cím vagy helyzetleírás")
        if [o["jel"] for o in s["opciok"]] != ["A", "B", "C", "D"]:
            hibak.append(f"{hol}: a válaszok jele {[o['jel'] for o in s['opciok']]}, A–D kellene")
        for o in s["opciok"]:
            for k in ("opcio_leirasa", "fogalom", "kiertekeles"):
                if not o[k]: hibak.append(f"{hol} {o['jel']}: üres mező ({k})")
    if [s["id"] for s in szituaciok] != list(range(1, 11)):
        hibak.append(f"{ws.title}: a helyzetek sorszáma {[s['id'] for s in szituaciok]}, 1–10 kellene")
    return dict(regi, szituaciok=szituaciok)

def main():
    ut = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else GYOKER / "szerkesztes" / "kerdesek.xlsx"
    wb = load_workbook(ut, data_only=True)
    tort = {}
    if "Történetek" in wb.sheetnames:
        for sor in wb["Történetek"].iter_rows(min_row=2, values_only=True):
            if sor and sor[0]: tort[szoveg(sor[0])] = (szoveg(sor[1]), szoveg(sor[2]))
    uj = {}
    for kulcs, nev in PALYAK:
        regi = json.loads((GYOKER / "adatok" / f"{kulcs}.json").read_text(encoding="utf-8"))
        if nev not in wb.sheetnames:
            hibak.append(f"hiányzó munkalap: {nev}"); continue
        if nev in tort:
            regi["szerep"], regi["tortenet_szinopszisa"] = tort[nev]
        uj[kulcs] = palya(wb[nev], regi)
    if hibak:
        print("Nem írtam semmit, mert a táblázatban hiba van:")
        for h in hibak: print("  •", h)
        sys.exit(1)
    for kulcs, d in uj.items():
        cel = GYOKER / "adatok" / f"{kulcs}.json"
        regi = cel.read_text(encoding="utf-8")
        friss = json.dumps(d, ensure_ascii=False, indent=2) + "\n"
        valtozott = sum(1 for a, b in zip(regi.splitlines(), friss.splitlines()) if a != b) + abs(len(regi.splitlines()) - len(friss.splitlines()))
        cel.write_text(friss, encoding="utf-8")
        print(f"{cel.relative_to(GYOKER)}: {valtozott} sor változott")

if __name__ == "__main__":
    main()
