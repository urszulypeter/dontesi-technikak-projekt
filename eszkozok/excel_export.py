#!/usr/bin/env python3
"""A kérdésadatbázis (adatok/*.json) kiírása szerkeszthető Excel-fájlba.

Futtatás a projekt gyökeréből:  python3 eszkozok/excel_export.py
Eredmény: szerkesztes/kerdesek.xlsx  — visszaolvasás: eszkozok/excel_import.py"""
import json, math, pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

GYOKER = pathlib.Path(__file__).resolve().parent.parent
KIMENET = GYOKER / "szerkesztes" / "kerdesek.xlsx"

PALYAK = [("politikai", "Politikai"), ("vallalati", "Vállalati"), ("katonai", "Katonai")]
PONTKULCSOK = ["kovetkezetesseg", "atlathatosag", "dontesi_minoseg", "felelossegvallalas", "erintettek_bevonasa"]
PONTFEJ = ["Következetesség", "Átláthatóság", "Döntési minőség", "Felelősségvállalás", "Érintettek bevonása"]
SULY = [0.22, 0.22, 0.20, 0.18, 0.18]
SZORZO = 3

# Oszlopok a pályalapokon. Az importáló a fejléc szövege alapján keres, nem betűjel alapján.
OSZLOPOK = [  # (fejléc, szélesség, szerkeszthető)
    ("Sz.", 5, False), ("Jelenet címe", 18, True), ("Helyzet (a játékos ezt olvassa)", 46, True),
    ("Jel", 5, False), ("Válasz (a játékos ezt olvassa)", 44, True), ("Fogalom (a jegyzetből)", 22, True),
    ("Kiértékelés (a válasz után jelenik meg)", 58, True),
] + [(f, 11, True) for f in PONTFEJ] + [
    ("Súlyozott hatás a Bizalomindexre", 13, False), ("Ellenőrzés", 18, False),
]

BETU = "Arial"
ZOLD_SOTET, KREM, ARANY = "13291F", "F6F1E7", "B8933F"
SZERK = PatternFill("solid", fgColor="FFFFFF")
ZART = PatternFill("solid", fgColor="EDEAE3")
FEJ = PatternFill("solid", fgColor=ZOLD_SOTET)
VONAL = Side(style="thin", color="C9C2B3")
KERET = Border(left=VONAL, right=VONAL, top=VONAL, bottom=VONAL)
VASTAG = Border(left=VONAL, right=VONAL, top=Side(style="medium", color=ZOLD_SOTET), bottom=VONAL)

def sorok(szoveg, szelesseg):
    """Becsült sorszám tördelt szövegnél (Arial 10, ~1,15 karakter / szélességegység)."""
    kar = max(1, int(szelesseg * 1.15))
    return sum(max(1, math.ceil(len(b) / kar)) for b in str(szoveg).split("\n"))

def fejlec(ws, sor, ertekek):
    for i, v in enumerate(ertekek, 1):
        c = ws.cell(sor, i, v)
        c.font = Font(name=BETU, bold=True, color="FFFFFF", size=10)
        c.fill = FEJ; c.border = KERET
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

def utmutato(wb, minta):
    ws = wb.active; ws.title = "Útmutató"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    sorok_ = [
        ("cim", "Vezetői Bizalom — kérdésadatbázis"),
        ("al", "Ebben a fájlban a játék mindhárom története szerkeszthető. Ha kész, küldd vissza, és visszatöltöm a játékba."),
        None,
        ("fej", "Mit hol találsz"),
        ("sor", "Történetek", "A három pálya neve és a történet rövid összefoglalója, amely a szerepválasztó kártyán jelenik meg."),
        ("sor", "Politikai / Vállalati / Katonai", "Pályánként 10 helyzet, mindegyikhez 4 válasz — egy sor egy válasz."),
        None,
        ("fej", "Mit szabad átírni"),
        ("sor", "Fehér cellák", "Szabadon átírhatók: jelenet címe, helyzetleírás, válasz, fogalom, kiértékelés és az öt pontszám."),
        ("sor", "Szürke cellák", "Azonosítók és számított oszlopok (Sz., Jel, Súlyozott hatás, Ellenőrzés). Ezeket ne írd át — a visszatöltés ezek alapján párosít."),
        ("sor", "Sorrend", "Sort ne szúrj be és ne törölj: minden helyzetnek pontosan 4 válasza van (A–D). A válaszok sorrendje a játékban úgyis véletlenszerű."),
        None,
        ("fej", "A pontok"),
        ("sor", "Tartomány", "Egész szám −3 és +3 között, mind az öt oszlopban (0 = nem érinti). A játék háromszorozza: +2 a Bizalomindex adott dimenziójában +6 pontot jelent."),
        ("sor", "Szabály", "Minden válasznak legyen legalább egy pozitív és legalább egy negatív pontja: nincs ingyen jó döntés. Az „Ellenőrzés” oszlop jelez, ha ez sérül."),
        ("sor", "Súlyozott hatás", "Mennyit mozdít a válasz a Bizalomindexen (0,22·Köv + 0,22·Átl + 0,20·Dönt + 0,18·Fel + 0,18·Bev, háromszorozva). Segít látni, melyik válasz a „legdrágább”."),
        ("sor", "Mi alapján?", "A pontozás szempontjait a külön kiküldött „Pontozási szempontok” dokumentum írja le, dimenziónként."),
        None,
        ("fej", "Nyelvi szabály"),
        ("sor", "Helyzet és válasz", "Hétköznapi nyelv, szakszó nélkül — a játékosnak nem kell ismernie a tananyagot."),
        ("sor", "Kiértékelés", "Itt kell megnevezni a jegyzet fogalmát (PrOACT, korlátozott racionalitás, Mintzberg szerepei, halo-hatás stb.), egy-két érthető mondatban."),
        None,
        ("fej", "Minta: egy kitöltött sor"),
        ("sor", "Válasz", minta["opcio_leirasa"]),
        ("sor", "Fogalom", minta["fogalom"]),
        ("sor", "Kiértékelés", minta["kiertekeles"]),
        ("sor", "Pontok", " · ".join(f"{f} {minta['pontok'][k]:+d}".replace("-", "−") for f, k in zip(PONTFEJ, PONTKULCSOK))),
    ]
    r = 2
    for s in sorok_:
        if s is None:
            r += 1; continue
        if s[0] == "cim":
            c = ws.cell(r, 2, s[1]); c.font = Font(name=BETU, bold=True, size=16, color=ZOLD_SOTET)
        elif s[0] == "al":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            c = ws.cell(r, 2, s[1]); c.font = Font(name=BETU, size=10, color="5A5A5A")
        elif s[0] == "fej":
            c = ws.cell(r, 2, s[1]); c.font = Font(name=BETU, bold=True, size=11, color=ARANY)
        else:
            a = ws.cell(r, 2, s[1]); a.font = Font(name=BETU, bold=True, size=10)
            b = ws.cell(r, 3, s[2]); b.font = Font(name=BETU, size=10)
            for c in (a, b): c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * sorok(s[2], 92) + 3
        r += 1

def tortenetek(wb, adatok):
    ws = wb.create_sheet("Történetek")
    ws.freeze_panes = "A2"
    for col, (f, w) in enumerate([("Pálya", 14), ("Szerep (kártya címe)", 24), ("A történet összefoglalója", 100)], 1):
        ws.column_dimensions["ABC"[col - 1]].width = w
    fejlec(ws, 1, ["Pálya", "Szerep (kártya címe)", "A történet összefoglalója"])
    ws.row_dimensions[1].height = 22
    for i, (kulcs, nev) in enumerate(PALYAK, 2):
        d = adatok[kulcs]
        for col, v in enumerate([nev, d["szerep"], d["tortenet_szinopszisa"]], 1):
            c = ws.cell(i, col, v)
            c.font = Font(name=BETU, size=10, bold=(col == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = KERET
            c.fill = ZART if col == 1 else SZERK
        ws.row_dimensions[i].height = 15 * sorok(d["tortenet_szinopszisa"], 100) + 6

def palya(wb, nev, d):
    ws = wb.create_sheet(nev)
    ws.freeze_panes = "E2"
    fejlec(ws, 1, [o[0] for o in OSZLOPOK])
    ws.row_dimensions[1].height = 44
    for i, (_, w, _) in enumerate(OSZLOPOK, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    pont = DataValidation(type="whole", operator="between", formula1="-3", formula2="3", allow_blank=False,
                          showErrorMessage=True, errorTitle="Érvénytelen pont",
                          error="Egész számot adj meg −3 és +3 között.")
    ws.add_data_validation(pont)

    r = 2
    for sz in d["szituaciok"]:
        elso = r
        for k, o in enumerate(sz["opciok"]):
            ertek = [sz["id"], sz["cim"] if k == 0 else None, sz["szituacio_leirasa"] if k == 0 else None,
                     o["jel"], o["opcio_leirasa"], o["fogalom"], o["kiertekeles"]] + [o["pontok"][p] for p in PONTKULCSOK]
            for col, v in enumerate(ertek, 1):
                if v is None: continue
                ws.cell(r, col, v)
            # Súlyozott hatás és ellenőrzés: képlet, hogy szerkesztés után is számoljon.
            h, i_, j, kk, l = (ws.cell(r, c).coordinate for c in range(8, 13))
            ws.cell(r, 13, f"=ROUND({SZORZO}*({SULY[0]}*{h}+{SULY[1]}*{i_}+{SULY[2]}*{j}+{SULY[3]}*{kk}+{SULY[4]}*{l}),1)")
            ws.cell(r, 14, f'=IF(COUNT({h}:{l})<5,"hiányzó pont",IF(OR(MIN({h}:{l})<-3,MAX({h}:{l})>3),"−3…+3 közé",'
                           f'IF(COUNTIF({h}:{l},">0")=0,"kell pozitív is",IF(COUNTIF({h}:{l},"<0")=0,"kell negatív is","rendben"))))')
            pont.add(f"{h}:{l}")

            magas = max(sorok(o["opcio_leirasa"], 44), sorok(o["kiertekeles"], 58), sorok(o["fogalom"], 22), 2)
            ws.row_dimensions[r].height = 14 * magas + 6
            r += 1
        # A cím és a helyzetleírás a négy válasz sorára összevonva.
        for col in (1, 2, 3):
            ws.merge_cells(start_row=elso, start_column=col, end_row=r - 1, end_column=col)
        szukseges = 14 * sorok(sz["szituacio_leirasa"], 46) + 10
        meglevo = sum(ws.row_dimensions[x].height for x in range(elso, r))
        if meglevo < szukseges:
            for x in range(elso, r):
                ws.row_dimensions[x].height += (szukseges - meglevo) / 4

        for x in range(elso, r):
            for col in range(1, len(OSZLOPOK) + 1):
                c = ws.cell(x, col)
                szerk = OSZLOPOK[col - 1][2]
                c.font = Font(name=BETU, size=10, bold=col in (1, 2, 4))
                c.fill = SZERK if szerk else ZART
                c.border = VASTAG if x == elso else KERET
                kozep = col in (1, 4) or col >= 8
                c.alignment = Alignment(wrap_text=True, vertical="top" if not kozep else "center",
                                        horizontal="center" if kozep else "left")
            ws.cell(x, 13).number_format = '+0.0;"−"0.0;0'

    utolso = r - 1
    pontok = f"H2:L{utolso}"
    ws.conditional_formatting.add(pontok, CellIsRule(operator="greaterThan", formula=["0"],
                                  font=Font(name=BETU, color="1E6B3A", bold=True), fill=PatternFill("solid", fgColor="E3F1E6")))
    ws.conditional_formatting.add(pontok, CellIsRule(operator="lessThan", formula=["0"],
                                  font=Font(name=BETU, color="9B2C2C", bold=True), fill=PatternFill("solid", fgColor="F8E3E1")))
    ws.conditional_formatting.add(f"N2:N{utolso}", FormulaRule(formula=[f'N2<>"rendben"'],
                                  font=Font(name=BETU, color="9B2C2C", bold=True), fill=PatternFill("solid", fgColor="F8E3E1")))
    ws.auto_filter.ref = f"A1:N{utolso}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

def main():
    adatok = {k: json.loads((GYOKER / "adatok" / f"{k}.json").read_text(encoding="utf-8")) for k, _ in PALYAK}
    wb = Workbook()
    minta = next(o for o in adatok["politikai"]["szituaciok"][0]["opciok"] if o["fogalom"].startswith("PrOACT"))
    utmutato(wb, minta)
    tortenetek(wb, adatok)
    for kulcs, nev in PALYAK:
        palya(wb, nev, adatok[kulcs])
    # A képleteknek nincs tárolt értékük — az Excel megnyitáskor számolja őket.
    wb.calculation.fullCalcOnLoad = True
    KIMENET.parent.mkdir(exist_ok=True)
    wb.save(KIMENET)
    print(f"Kész: {KIMENET.relative_to(GYOKER)}")

if __name__ == "__main__":
    main()
